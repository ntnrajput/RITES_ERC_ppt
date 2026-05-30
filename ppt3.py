import pandas as pd
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
import os
import copy
from difflib import get_close_matches
import sys

# ================== CONFIG ==================

excel_file = sys.argv[1]
sheet_name = "Copy of Final Month Wise Defect"

template_ppt = "RITES_Template.pptx"
output_ppt = "Defect_Percentage_Pie_RITES_Final.pptx"

# ================== TARGET DEFECTS ==================
TARGET_DEFECTS = [
    "Application & Deflection test",
    "Chemical Composition",
    "Toe load",
    "Hardness",
    "Microstructure",
    "Grain size",
    "Depth of Decarburization",
    "Inclusion rating",
    "Marking",
    "PIO Rejection Reason"
]

# ================== PIO RED COLUMNS ==================
# ONLY COUNT COLUMNS (NO % COLUMNS)

PIO_COLUMNS = [
    "MPI",
    "Forging defect",
    "Embossing Die Impression Failure",
    "Other (less weigh)"
]

# ================== FIXED COLORS ==================
DEFECT_COLORS = {
    "Application & Deflection test": "#1f4e79",
    "Chemical Composition": "#4f81bd",
    "Toe load": "#9bbb59",
    "Hardness": "#f79646",
    "Microstructure": "#8064a2",
    "Grain size": "#c0504d",
    "Depth of Decarburization": "#00b0f0",
    "Inclusion rating": "#7030a0",
    "Marking": "#ffc000",
    "PIO Rejection Reason": "#ff0000"
}

# ================== LOAD EXCEL ==================
df = pd.read_excel(excel_file, sheet_name=sheet_name, header=1)

# CLEAN COLUMN NAMES
df.columns = (
    df.columns.astype(str)
    .str.strip()
    .str.replace("\n", " ", regex=True)
)

# RENAME PLANT COLUMN
df.rename(columns={"Firm Name": "Plant"}, inplace=True)

# ================== MAP DEFECT COLUMNS ==================
def map_defect_columns(df, target_defects):

    mapped = {}

    for target in target_defects:

        if target == "PIO Rejection Reason":
            continue

        match = get_close_matches(target, df.columns, n=1, cutoff=0.6)

        if match:
            mapped[target] = match[0]

    return mapped

defect_map = map_defect_columns(df, TARGET_DEFECTS)

# ================== CALCULATE DEFECT DATA ==================
def calculate_defect_data(df, defect_map):

    results = []

    for plant, group in df.groupby("Plant"):

        inspected_total = group["Inspected (Nos.)"].sum()

        if inspected_total == 0:
            continue

        row_pct = {"Plant": plant}
        row_count = {}

        # -------- NORMAL DEFECTS --------
        for defect_name, col in defect_map.items():

            rejected_total = group[col].sum()

            pct = (rejected_total / inspected_total) * 100

            row_pct[defect_name] = pct
            row_count[defect_name] = rejected_total

        # -------- PIO REJECTION TOTAL --------
        pio_total = 0

        for col in PIO_COLUMNS:

            matched = get_close_matches(col, group.columns, n=1, cutoff=0.6)

            if matched:
                pio_total += group[matched[0]].sum()

        pio_pct = (pio_total / inspected_total) * 100

        row_pct["PIO Rejection Reason"] = pio_pct
        row_count["PIO Rejection Reason"] = pio_total

        results.append((row_pct, row_count))

    return results

data_combined = calculate_defect_data(df, defect_map)

# ================== PIE CHART ==================
def plot_pie(row_pct, row_count, img_path):

    labels = []
    sizes = []
    counts = []

    for defect in TARGET_DEFECTS:

        pct = row_pct.get(defect, 0)
        cnt = row_count.get(defect, 0)

        if cnt > 0:
            labels.append(defect)
            sizes.append(pct)
            counts.append(cnt)

    if not sizes:
        return False

    colors = [DEFECT_COLORS[d] for d in labels]

    def autopct_func(actual_pcts, counts):

        counter = {'i': 0}

        def inner(_pct):

            i = counter['i']

            actual_pct = actual_pcts[i]
            count = counts[i]

            counter['i'] += 1

            return f"{int(count)} ({actual_pct:.2f}%)"

        return inner

    plt.figure(figsize=(5, 5))

    plt.pie(
        sizes,
        labels=None,
        colors=colors,
        autopct=autopct_func(sizes, counts),
        startangle=140
    )

    plt.axis("equal")

    plt.savefig(img_path, dpi=300, bbox_inches="tight")

    plt.close()

    return True

# ================== PPT HELPERS ==================
def duplicate_slide(prs, slide):

    new_slide = prs.slides.add_slide(slide.slide_layout)

    for shape in slide.shapes:

        el = shape.element

        new_slide.shapes._spTree.insert_element_before(
            copy.deepcopy(el),
            'p:extLst'
        )

    return new_slide

def clear_old_charts(slide):

    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in list(slide.shapes):

        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            slide.shapes._spTree.remove(shape.element)

def add_title(slide, text):

    box = slide.shapes.add_textbox(
        Inches(0.5),
        Inches(0.3),
        Inches(11),
        Inches(0.5)
    )

    p = box.text_frame.paragraphs[0]

    p.text = text
    p.font.size = Pt(20)
    p.font.bold = True
    p.alignment = 1

def add_plant_title(slide, text, left):

    box = slide.shapes.add_textbox(
        left,
        Inches(0.8),
        Inches(4.8),
        Inches(0.4)
    )

    p = box.text_frame.paragraphs[0]

    p.text = text
    p.font.size = Pt(14)
    p.font.bold = True
    p.alignment = 1

# ================== DYNAMIC LEGEND ==================
def add_dynamic_legend(slide, legend_defects):

    start_left = Inches(10.5)
    start_top = Inches(1.5)

    for i, defect in enumerate(legend_defects):

        color = DEFECT_COLORS.get(defect, "#cccccc")

        top = start_top + Inches(i * 0.4)

        # Color Box
        rect = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.RECTANGLE,
            start_left,
            top,
            Inches(0.25),
            Inches(0.25)
        )

        rect.fill.solid()

        rect.fill.fore_color.rgb = RGBColor.from_string(
            color.replace("#", "")
        )

        rect.line.fill.background()

        # Text
        textbox = slide.shapes.add_textbox(
            start_left + Inches(0.35),
            top - Inches(0.05),
            Inches(3),
            Inches(0.3)
        )

        p = textbox.text_frame.paragraphs[0]

        p.text = defect
        p.font.size = Pt(10)

# ================== BUILD PPT ==================
prs = Presentation(template_ppt)

base_slide = prs.slides[0]

os.makedirs("charts_pct", exist_ok=True)

for i in range(0, len(data_combined), 2):

    slide = duplicate_slide(prs, base_slide)

    clear_old_charts(slide)

    add_title(
        slide,
        "Defect % Analysis (Against Inspected Quantity)"
    )

    # -------- COLLECT LEGEND DEFECTS --------
    legend_defects = []

    for k in range(i, min(i + 2, len(data_combined))):

        row_pct, row_count = data_combined[k]

        for defect, count in row_count.items():

            if count > 0 and defect not in legend_defects:
                legend_defects.append(defect)

    # -------- ADD LEGEND --------
    add_dynamic_legend(slide, legend_defects)

    # -------- ADD CHARTS --------
    for j in range(2):

        if i + j >= len(data_combined):
            break

        row_pct, row_count = data_combined[i + j]

        plant = row_pct["Plant"]

        safe_plant_name = (
            plant.replace("/", "_")
                 .replace("\\", "_")
        )

        img_path = f"charts_pct/{safe_plant_name}.png"

        success = plot_pie(
            row_pct,
            row_count,
            img_path
        )

        if not success:
            continue

        left = Inches(1) if j == 0 else Inches(6)

        add_plant_title(slide, plant, left)

        slide.shapes.add_picture(
            img_path,
            left,
            Inches(1.3),
            width=Inches(4.5),
            height=Inches(3.5)
        )

# ================== REMOVE TEMPLATE SLIDE ==================
prs.slides._sldIdLst.remove(
    prs.slides._sldIdLst[0]
)

# ================== SAVE PPT ==================
prs.save(output_ppt)

print(" FINAL PPT CREATED:", output_ppt)