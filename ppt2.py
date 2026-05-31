import pandas as pd
import matplotlib.pyplot as plt
import re
import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE_TYPE
import copy

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

template_ppt = os.path.join(
    BASE_DIR,
    "RITES_Template.pptx"
)

OUTPUT_DIR = os.path.join(BASE_DIR, "output")
os.makedirs(OUTPUT_DIR, exist_ok=True)

output_ppt = os.path.join(
    OUTPUT_DIR,
    "Vendor_PO_Quality_Analysis.pptx"
)

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

prs = Presentation(template_ppt)
base_slide = prs.slides[0]


def duplicate_slide(prs, slide):
    new_slide = prs.slides.add_slide(slide.slide_layout)
    for shape in slide.shapes:
        el = shape.element
        new_slide.shapes._spTree.insert_element_before(
            copy.deepcopy(el),
            'p:extLst'
        )
    return new_slide


def clear_old_charts(slide):
    for shape in list(slide.shapes):
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            slide.shapes._spTree.remove(shape.element)

# =========================
# 1️⃣ LOAD DATA
# =========================
excel_file = sys.argv[1]
sheet_name = "Plant, Month & PO Defect Analys"

df = pd.read_excel(excel_file, sheet_name=sheet_name, header=1)
df.columns = df.columns.str.strip()

df.rename(columns={
    "Firm Name": "Vendor",
    "PO Number /Date": "PO",
}, inplace=True)


def normalize_po_value(value):
    if pd.isna(value):
        return pd.NA

    if isinstance(value, str):
        value = re.sub(r"\s+", " ", value).strip()
        return value if value else pd.NA

    if isinstance(value, (int, float)):
        if float(value).is_integer():
            return str(int(value))
        return str(value).strip()

    value = str(value).strip()
    return value if value else pd.NA


def prepare_po_chart_data(df_vendor):
    df_vendor = df_vendor.copy()
    df_vendor["PO"] = df_vendor["PO"].apply(normalize_po_value)
    df_vendor = df_vendor[df_vendor["PO"].notna()]
    df_vendor["PO_Label"] = df_vendor["PO"].astype(str)
    return df_vendor.sort_values("PO_Label", kind="stable").reset_index(drop=True)


def build_chart_path(folder, base_name, max_length=40):

    folder = os.path.join(BASE_DIR, folder)

    os.makedirs(folder, exist_ok=True)

    safe_name = re.sub(r'[<>:"/\\\\|?*]', "_", str(base_name))
    safe_name = re.sub(r"\s+", " ", safe_name).strip(" .")
    safe_name = safe_name[:max_length].rstrip(" .") or "chart"

    file_path = os.path.join(folder, f"{safe_name}.png")

    if os.path.exists(file_path):
        os.remove(file_path)

    return file_path


df["PO"] = df["PO"].apply(normalize_po_value)

# ================== VENDOR–PO SUMMARY ==================

df_po = df.copy()

df_po.rename(columns={
    "Plant": "Vendor",
    "PO Number /Date": "PO"
}, inplace=True)

po_summary = (
    df_po.groupby(["Vendor", "PO"], as_index=False)
    .agg({
        "Inspected (Nos.)": "sum",
        "Acceptance (Nos.)": "sum",
        "Rejection (Nos.)": "sum"
    })
)

po_summary["Rejection %"] = (
    po_summary["Rejection (Nos.)"] /
    po_summary["Inspected (Nos.)"] * 100
).fillna(0)

vendor_po_count = (
    po_summary.groupby("Vendor")["PO"]
    .nunique()
    .reset_index(name="No. of POs")
    .sort_values("No. of POs", ascending=False)
)



# =========================
# 2️⃣ CLEAN VENDOR NAME
# =========================
rio_prefixes = ["NRIO", "SRIO", "ERIO", "WRIO", "CRIO"]

def plot_vendor_po_count(df, img_path):

    plt.figure(figsize=(11, 6))
    plt.barh(df["Vendor"], df["No. of POs"], color="#4f81bd")
    plt.xlabel("Number of Running POs")
    plt.title("Vendor-wise Number of Running POs")
    plt.gca().invert_yaxis()
    plt.tight_layout()
    plt.savefig(img_path, dpi=300)
    plt.close()
def plot_vendor_po_count(df, img_path):

    plt.figure(figsize=(11, 6))
    plt.barh(df["Vendor"], df["No. of POs"], color="#4f81bd")
    plt.xlabel("Number of Running POs")
    plt.title("Vendor-wise Number of Running POs")
    plt.gca().invert_yaxis()
    plt.tight_layout()
    plt.savefig(img_path, dpi=300)
    plt.close()


def plot_zone_po_count(df, img_path):

    plt.figure(figsize=(11, 6))

    plt.bar(
        df["Zonal Railway (Purchaser)"],
        df["No. of POs"]
    )

    plt.xticks(rotation=45, ha="right")
    plt.xlabel("Zonal Railway")
    plt.ylabel("Number of POs")
    plt.title("Zone-wise Number of Purchase Orders")

    plt.tight_layout()
    plt.savefig(img_path, dpi=300)
    plt.close()


def add_zone_po_slide(prs, base_slide, zone_po_count):

    img_path = build_chart_path("po_charts", "zone_po_count")

    plot_zone_po_count(zone_po_count, img_path)

    slide = duplicate_slide(prs, base_slide)
    clear_old_charts(slide)

    title = slide.shapes.add_textbox(
        Inches(0.5), Inches(0.3), Inches(11), Inches(0.6)
    )

    p = title.text_frame.paragraphs[0]
    p.text = "Zone-wise Number of Purchase Orders"
    p.font.size = Pt(20)
    p.font.bold = True
    p.alignment = 1

    slide.shapes.add_picture(
        img_path,
        Inches(0.5),
        Inches(1.1),
        width=Inches(11)
    )




def plot_po_quality_vendor(vendor, df_vendor, img_path):

    df_vendor = prepare_po_chart_data(df_vendor)
    if df_vendor.empty:
        return

    x = range(len(df_vendor))

    fig, ax1 = plt.subplots(figsize=(11, 5))

    ax1.bar(
        x,
        df_vendor["Acceptance (Nos.)"],
        label="Accepted",
        color="#9bbb59"
    )

    ax1.bar(
        x,
        df_vendor["Rejection (Nos.)"],
        bottom=df_vendor["Acceptance (Nos.)"],
        label="Rejected",
        color="#c0504d"
    )

    ax1.set_ylabel("Quantity (Nos.)")
    ax1.set_xticks(x)
    ax1.set_xticklabels(df_vendor["PO_Label"], rotation=45, ha="right", fontsize=8)

    ax2 = ax1.twinx()
    ax2.plot(
        x,
        df_vendor["Rejection %"],
        color="#1f4e79",
        marker="o",
        linewidth=2
    )
    ax2.set_ylabel("Rejection (%)")
    ax2.set_ylim(0, max(1, df_vendor["Rejection %"].max() * 1.3))

    ax1.legend(loc="upper left")
    plt.title(f"PO-wise Inspection Performance – {vendor}", fontsize=14, fontweight="bold")

    plt.tight_layout()
    plt.savefig(img_path, dpi=300)
    plt.close()


def add_vendor_po_load_slide(prs, base_slide, vendor_po_count):

    img_path = build_chart_path("po_charts", "vendor_po_count")
    plot_vendor_po_count(vendor_po_count, img_path)

    slide = duplicate_slide(prs, base_slide)
    clear_old_charts(slide)

    title = slide.shapes.add_textbox(
        Inches(0.5), Inches(0.3), Inches(11), Inches(0.6)
    )
    p = title.text_frame.paragraphs[0]
    p.text = "Vendor-wise Running Purchase Orders"
    p.font.size = Pt(20)
    p.font.bold = True
    p.alignment = 1

    slide.shapes.add_picture(
        img_path,
        Inches(0.5),
        Inches(1.1),
        width=Inches(11)
    )


def add_vendor_po_quality_slides(prs, base_slide, po_summary, vendor_po_count, top_n=6):

    top_vendors = vendor_po_count.head(top_n)["Vendor"]

    for vendor in top_vendors:

        vendor_df = po_summary[po_summary["Vendor"] == vendor]
        img_path = build_chart_path("po_charts", f"po_quality_{vendor[:30]}")

        plot_po_quality_vendor(vendor, vendor_df, img_path)

        slide = duplicate_slide(prs, base_slide)
        clear_old_charts(slide)

        title = slide.shapes.add_textbox(
            Inches(0.5), Inches(0.3), Inches(11), Inches(0.6)
        )
        p = title.text_frame.paragraphs[0]
        p.text = f"PO-wise Quality Performance – {vendor}"
        p.font.size = Pt(18)
        p.font.bold = True
        p.alignment = 1

        slide.shapes.add_picture(
            img_path,
            Inches(0.5),
            Inches(1.1),
            width=Inches(11)
        )



def clean_vendor(name):
    if not isinstance(name, str):
        return name
    name = name.strip()
    for p in rio_prefixes:
        if name.startswith(p):
            name = name.replace(p, "", 1).strip()
    # remove extra spaces
    name = re.sub(r"\s+", " ", name)
    return name

df["Vendor"] = df["Vendor"].apply(clean_vendor)

# =========================
# 3️⃣ PO-WISE SUMMARY
# =========================
po_summary = (
    df.groupby(["Vendor", "PO"], as_index=False)
    .agg({
        "Inspected (Nos.)": "sum",
        "Acceptance (Nos.)": "sum",
        "Rejection (Nos.)": "sum"
    })
)

po_summary["Rejection %"] = (
    po_summary["Rejection (Nos.)"] /
    po_summary["Inspected (Nos.)"] * 100
).fillna(0)

# =========================
# 4️⃣ NO. OF POs PER VENDOR
# =========================
vendor_po_count = (
    po_summary.groupby("Vendor")["PO"]
    .nunique()
    .reset_index(name="No. of POs")
    .sort_values("No. of POs", ascending=False)
)

# =========================
# 9️⃣ ZONE-WISE PO COUNT
# =========================

df_zone = df.copy()

# Clean column name (remove extra spaces)
df_zone.columns = df_zone.columns.str.strip()

# Remove NO PRODUCTION / NA POs
df_zone = df_zone[
    df_zone["PO"].notna() &
    (df_zone["PO"].astype(str).str.upper() != "NO PRODUCTION") &
    (df_zone["PO"].astype(str).str.upper() != "NA")
]

zone_po_count = (
    df_zone.groupby("Zonal Railway (Purchaser)")["PO"]
    .nunique()
    .reset_index(name="No. of POs")
    .sort_values("No. of POs", ascending=False)
)

print("✅ Zone-wise PO summary created")

# =========================
# 5️⃣ CHART 1 – VENDOR vs NO. OF POs
# =========================
plt.figure(figsize=(10, 6))

plt.barh(
    vendor_po_count["Vendor"],
    vendor_po_count["No. of POs"]
)

plt.xlabel("Number of Running POs")
plt.title("Vendor-wise Number of Running POs")
plt.gca().invert_yaxis()

plt.tight_layout()
plt.savefig("vendor_po_count.png", dpi=300)
plt.close()

print("✅ Chart saved: vendor_po_count.png")

# =========================
# 6️⃣ CHART 2 – PO-WISE QUALITY (PER VENDOR)
# =========================

PO_CHART_DIR = os.path.join(BASE_DIR, "po_charts")
os.makedirs(PO_CHART_DIR, exist_ok=True)

def plot_po_quality(vendor_name, df_vendor):
    df_vendor = prepare_po_chart_data(df_vendor)
    if df_vendor.empty:
        print(f"Skipping PO chart for {vendor_name}: no valid PO values found")
        return

    x = range(len(df_vendor))

    fig, ax1 = plt.subplots(figsize=(11, 5))

    # Accepted & Rejected bars
    ax1.bar(
        x,
        df_vendor["Acceptance (Nos.)"],
        label="Accepted",
        alpha=0.85
    )

    ax1.bar(
        x,
        df_vendor["Rejection (Nos.)"],
        bottom=df_vendor["Acceptance (Nos.)"],
        label="Rejected"
    )

    ax1.set_ylabel("Quantity (Nos.)")
    ax1.set_xticks(x)
    ax1.set_xticklabels(df_vendor["PO_Label"], rotation=45, ha="right", fontsize=8)

    # Rejection % line
    ax2 = ax1.twinx()
    ax2.plot(
        x,
        df_vendor["Rejection %"],
        marker="o",
        linewidth=2
    )
    ax2.set_ylabel("Rejection (%)")
    ax2.set_ylim(0, max(df_vendor["Rejection %"].max() * 1.3, 1))

    ax1.set_title(f"PO-wise Inspection Performance – {vendor_name}")
    ax1.legend(loc="upper left")

    plt.tight_layout()
    file_name = build_chart_path("po_charts", vendor_name)
    plt.savefig(file_name, dpi=300)
    plt.close()

    print(f"✅ PO chart saved for {vendor_name}")

# =========================
# 7️⃣ GENERATE PO-WISE CHARTS (TOP VENDORS)
# =========================
TOP_N_VENDORS = 10

top_vendors = vendor_po_count.head(TOP_N_VENDORS)["Vendor"]

for vendor in top_vendors:
    vendor_df = po_summary[po_summary["Vendor"] == vendor]
    plot_po_quality(vendor, vendor_df)

# =========================
# 8️⃣ OPTIONAL: EXPORT SUMMARY TABLE
# =========================
po_summary.to_excel("vendor_po_wise_quality.xlsx", index=False)

add_vendor_po_load_slide(prs, base_slide, vendor_po_count)
add_vendor_po_quality_slides(prs, base_slide, po_summary, vendor_po_count)
add_zone_po_slide(prs, base_slide, zone_po_count)
prs.slides._sldIdLst.remove(prs.slides._sldIdLst[0])  # remove template slide
prs.save(output_ppt)

print("✅ PPT created:", output_ppt)
print("Saved PPT:", output_ppt)
print("Exists:", os.path.exists(output_ppt))
print("\n🎉 DONE!")
print("• Vendor-wise PO count chart")
print("• PO-wise Accepted / Rejected + Rejection % charts")
print("• Excel summary exported")
